import hashlib
import json
from pathlib import Path

import pymupdf
import pytest

from local_ocr.common import AppError, Cancelled
from local_ocr.dependencies import HOCR_CONFIG, available_languages, ensure_ocr_configs, find_program
from local_ocr.documents import DocumentOptions, convert_document


def make_pdf(path, text="Native searchable text for Local OCR testing."):
    with pymupdf.open() as pdf:
        page = pdf.new_page(width=612, height=792)
        page.insert_text((60, 80), text, fontsize=16)
        pdf.save(path)


def make_scan(path, mixed=False):
    with pymupdf.open() as native:
        page = native.new_page(width=500, height=180)
        page.insert_text((35, 60), "SCANNED TEXT RECOGNITION TEST", fontsize=20)
        page.insert_text((35, 100), "This sentence lives inside an image.", fontsize=16)
        image = page.get_pixmap(dpi=200).tobytes("png")
    with pymupdf.open() as pdf:
        page = pdf.new_page(width=612, height=792)
        if mixed:
            page.insert_text((50, 60), "NATIVE HEADER MUST BE PRESERVED", fontsize=15)
        page.insert_image(pymupdf.Rect(40, 130, 540, 310), stream=image)
        pdf.save(path)
    return image


def require_tesseract():
    if not find_program("tesseract"):
        pytest.skip("Tesseract no instalado")
    if "eng" not in available_languages({}):
        pytest.skip("Falta idioma eng")


def test_private_tessdata_repairs_required_hocr_config(tmp_path):
    tessdata = tmp_path / "tessdata"
    tessdata.mkdir()
    ensure_ocr_configs({"tessdata": str(tessdata)})
    assert (tessdata / "configs" / "hocr").read_text(encoding="ascii") == HOCR_CONFIG


def test_private_tessdata_keeps_existing_hocr_config(tmp_path):
    hocr = tmp_path / "tessdata" / "configs" / "hocr"
    hocr.parent.mkdir(parents=True)
    hocr.write_text("custom config\n", encoding="ascii")
    ensure_ocr_configs({"tessdata": str(tmp_path / "tessdata")})
    assert hocr.read_text(encoding="ascii") == "custom config\n"


def test_native_pdf_passthrough_and_markdown(tmp_path):
    source = tmp_path / "native.pdf"
    make_pdf(source)
    original = source.read_bytes()
    result = convert_document(source, tmp_path / "out")
    assert Path(result["pdf"]).read_bytes() == original
    assert "Native searchable text" in Path(result["markdown"]).read_text()
    assert source.read_bytes() == original
    report = json.loads((Path(result["output"]) / "informe.json").read_text())
    assert not report["ocr_applied"]
    assert report["pages"][0]["characters"] > 0


@pytest.mark.integration
@pytest.mark.parametrize("mixed", [False, True])
def test_real_ocr_scan_and_mixed_preserve_visual(tmp_path, mixed):
    require_tesseract()
    source = tmp_path / "scan.pdf"
    make_scan(source, mixed)
    original_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    result = convert_document(source, tmp_path / "out", DocumentOptions(language="eng", images=True))
    md = Path(result["markdown"]).read_text(encoding="utf-8")
    assert "SCANNED TEXT RECOGNITION TEST" in md
    assert "inside an image" in md
    if mixed:
        assert "NATIVE HEADER MUST BE PRESERVED" in md
    with pymupdf.open(source) as before, pymupdf.open(result["pdf"]) as after:
        assert len(after) == len(before)
        # El texto añadido es invisible, sin alterar los píxeles renderizados.
        assert before[0].get_pixmap().samples == after[0].get_pixmap().samples
        assert "RECOGNITION" in after[0].get_text()
    assert hashlib.sha256(source.read_bytes()).hexdigest() == original_hash


@pytest.mark.integration
def test_word_with_embedded_scan_to_pdf_and_markdown(tmp_path):
    require_tesseract()
    if not find_program("soffice"):
        pytest.skip("LibreOffice no instalado")
    from docx import Document
    from docx.shared import Inches

    image = make_scan(tmp_path / "scan.pdf")
    png = tmp_path / "scan.png"
    png.write_bytes(image)
    word = Document()
    word.add_heading("Word conversion test", 0)
    word.add_paragraph("Native Word text remains selectable.")
    word.add_picture(str(png), width=Inches(5))
    source = tmp_path / "word.docx"
    word.save(source)
    result = convert_document(source, tmp_path / "out", DocumentOptions(language="eng"))
    md = Path(result["markdown"]).read_text(encoding="utf-8")
    assert "Native Word text remains selectable" in md
    assert "SCANNED TEXT RECOGNITION TEST" in md


def test_missing_mode_warns_for_mixed_page(tmp_path):
    source = tmp_path / "mixed.pdf"
    make_scan(source, mixed=True)
    result = convert_document(source, tmp_path / "out", DocumentOptions(mode="missing"))
    assert any("imágenes de páginas" in s for s in result["warnings"])
    assert Path(result["pdf"]).read_bytes() == source.read_bytes()


def test_encrypted_pdf_rejected(tmp_path):
    path = tmp_path / "secret.pdf"
    with pymupdf.open() as pdf:
        pdf.new_page()
        pdf.save(
            path, encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="example-owner", user_pw="example-user"
        )
    with pytest.raises(AppError, match="protegido"):
        convert_document(path, tmp_path / "out")
    assert not list((tmp_path / "out").iterdir())


def test_signature_field_rejected(tmp_path):
    path = tmp_path / "signature.pdf"
    with pymupdf.open() as pdf:
        page = pdf.new_page()
        widget = pymupdf.Widget()
        widget.field_type = pymupdf.PDF_WIDGET_TYPE_SIGNATURE
        widget.field_name = "signature"
        widget.rect = pymupdf.Rect(50, 50, 200, 100)
        page.add_widget(widget)
        pdf.save(path)
    with pytest.raises(AppError, match="firma"):
        convert_document(path, tmp_path / "out")


def test_missing_languages_roll_back(tmp_path):
    if not find_program("tesseract"):
        pytest.skip("Tesseract no instalado")
    source = tmp_path / "scan.pdf"
    make_scan(source)
    with pytest.raises(AppError, match="Faltan idiomas"):
        convert_document(source, tmp_path / "out", DocumentOptions(language="nonexistent"))
    assert not list((tmp_path / "out").iterdir())


def test_cancel_leaves_no_outputs(tmp_path):
    source = tmp_path / "native.pdf"
    make_pdf(source)
    stop = tmp_path / "cancel"
    stop.touch()
    with pytest.raises(Cancelled):
        convert_document(source, tmp_path / "out", cancel_file=stop)
    assert not (tmp_path / "out").exists()


def test_options_reject_incompatible_deskew():
    with pytest.raises(AppError, match="Enderezar"):
        DocumentOptions(deskew=True).validate()


def test_broken_pdf_has_no_partial_results(tmp_path):
    source = tmp_path / "broken.pdf"
    source.write_bytes(b"not a pdf")
    with pytest.raises(Exception):
        convert_document(source, tmp_path / "out")
    assert not list((tmp_path / "out").iterdir())


@pytest.mark.integration
def test_force_mode_ocr_native_text(tmp_path):
    require_tesseract()
    source = tmp_path / "native.pdf"
    make_pdf(source, "FORCED OCR SAMPLE TEXT")
    result = convert_document(
        source, tmp_path / "out", DocumentOptions(language="eng", mode="force", deskew=True)
    )
    assert "FORCED OCR SAMPLE TEXT" in Path(result["markdown"]).read_text(encoding="utf-8")
    assert any("rasterizaron" in w for w in result["warnings"])
